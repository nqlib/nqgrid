"""openpyxl writer against the golden ``planning.sheet.json``."""

from __future__ import annotations

import io
from typing import Any

import pytest

from nqgrid_export import write_xlsx_from_sheet
from nqgrid_export.xlsx import heat_fill_hex, heat_step, xlsx_sheet_name

openpyxl = pytest.importorskip("openpyxl")
from openpyxl.cell.cell import MergedCell  # noqa: E402


@pytest.fixture(scope="module")
def workbook(planning_sheet: dict[str, Any], planning_spec: dict[str, Any]) -> Any:
    return openpyxl.load_workbook(io.BytesIO(write_xlsx_from_sheet(planning_sheet, spec=planning_spec)))


def test_values_and_number_formats(workbook: Any, planning_sheet: dict[str, Any]) -> None:
    ws = workbook["FY26 Sales Plan"]
    cells = planning_sheet["cells"]
    assert ws.max_row == len(cells)
    assert ws.max_column == planning_sheet["columnCount"]
    for r, row in enumerate(cells, start=1):
        for c, src in enumerate(row, start=1):
            cell = ws.cell(row=r, column=c)
            if src["v"] is None:
                assert cell.value is None
            elif isinstance(src["v"], float):
                assert cell.value == pytest.approx(src["v"])
            else:
                assert cell.value == src["v"]
            if isinstance(cell, MergedCell):
                continue  # covered by a merge: Excel keeps only the anchor's value and style
            if src.get("numFmt"):
                assert cell.number_format == src["numFmt"], (r, c)
            assert bool(cell.font.bold) == bool(src.get("bold")), (r, c)


def test_number_formats_present(workbook: Any, planning_sheet: dict[str, Any]) -> None:
    expected = {c["numFmt"] for row in planning_sheet["cells"] for c in row if c.get("numFmt")}
    actual = {
        c.number_format for row in workbook.worksheets[0].iter_rows() for c in row if isinstance(c.value, int | float)
    }
    assert actual == expected
    assert "#,##0%" in actual
    assert any(f.startswith("[>=1000000]") for f in actual)  # compact currency


def test_merges(workbook: Any, planning_sheet: dict[str, Any]) -> None:
    ws = workbook.worksheets[0]
    expected = {
        (m["row"] + 1, m["col"] + 1, m["row"] + m["rowSpan"], m["col"] + m["colSpan"]) for m in planning_sheet["merges"]
    }
    actual = {(r.min_row, r.min_col, r.max_row, r.max_col) for r in ws.merged_cells.ranges}
    assert actual == expected
    assert len(actual) == 3


def test_frozen_panes(workbook: Any, planning_sheet: dict[str, Any]) -> None:
    assert (planning_sheet["frozenRows"], planning_sheet["frozenCols"]) == (1, 2)
    assert workbook.worksheets[0].freeze_panes == "C2"


def test_heat_levels_become_fills(workbook: Any, planning_sheet: dict[str, Any]) -> None:
    ws = workbook.worksheets[0]
    heated = 0
    for r, row in enumerate(planning_sheet["cells"], start=1):
        for c, src in enumerate(row, start=1):
            fill = ws.cell(row=r, column=c).fill
            if "heat" in src:
                heated += 1
                assert fill.fill_type == "solid"
                assert fill.start_color.rgb == "FF" + heat_fill_hex(src["heat"]).lstrip("#").upper()
    assert heated > 50


def test_hidden_spec_sheet(workbook: Any) -> None:
    ss = workbook["Spec"]
    assert ss.sheet_state == "hidden"
    assert '"id": "planning-fy26",' in [c.value.strip() for c in ss["A"] if isinstance(c.value, str)]


def test_heat_palette_matches_typescript() -> None:
    # pivotHeatLevelFillHex(level, 9, "light") from @nqlib/nqgrid/render-static.
    assert heat_step(0, 9) == 0
    assert heat_step(8, 9) == 8
    assert heat_step(2, 5) == 4
    assert [heat_fill_hex(level) for level in (0, 1, 4, 8)] == ["#fef5f0", "#fde9de", "#f8c0a3", "#f2854b"]
    assert heat_fill_hex(2, 5) == "#f8c0a3"
    assert heat_fill_hex(0, theme="dark") == "#261d18"
    assert heat_fill_hex(3, 7, "dark") == "#723a1c"


def test_formula_like_strings_stay_text() -> None:
    sheet = {
        "cells": [[{"v": '=HYPERLINK("http://x")', "role": "leaf"}]],
        "merges": [],
        "frozenRows": 0,
        "frozenCols": 0,
    }
    ws = openpyxl.load_workbook(io.BytesIO(write_xlsx_from_sheet(sheet))).worksheets[0]
    assert ws["A1"].data_type == "s"
    assert ws.freeze_panes is None


def test_sheet_names() -> None:
    assert xlsx_sheet_name("FY26 / Q1: [draft]") == "FY26 Q1 draft"
    assert xlsx_sheet_name("x" * 40) == "x" * 31
    assert xlsx_sheet_name(None) == "Pivot"
